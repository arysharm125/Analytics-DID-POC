"""
Example Backend for Excel DID Generator

This backend demonstrates the proper architecture for integrating with DIDSvc:
- Frontend calls this backend (not DIDSvc directly)
- Backend generates the Excel file
- Backend registers the artefact with DIDSvc
- Backend returns the file to the frontend

The frontend is served as a static HTML file from this same server.
"""

import hashlib
import io
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path

import base58
import httpx
import qrcode
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image

# Configuration
DIDSVC_BASE_URL = os.getenv("DIDSVC_BASE_URL", "http://localhost:8000")
DIDSVC_API_TOKEN = os.getenv("DIDSVC_API_TOKEN", "example-api-token-123")
DIDCHECK_URL = os.getenv("DIDCHECK_URL", "http://localhost:5173")
EXAMPLE_APP_URL = os.getenv("EXAMPLE_APP_URL", "http://localhost:3001")

# Path to static files
STATIC_DIR = Path(__file__).parent / "static"

app = FastAPI(
    title="Excel DID Generator Backend",
    description="Example backend that generates Excel reports and registers them with DIDSvc",
)


def generate_sample_data() -> list[dict]:
    """Generate sample report data."""
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    return [
        {"Name": "Report 1", "Value": 100, "Date": today},
        {"Name": "Report 2", "Value": 200, "Date": today},
        {"Name": "Report 3", "Value": 300, "Date": today},
        {"Name": "Report 4", "Value": 400, "Date": today},
        {"Name": "Report 5", "Value": 500, "Date": today},
    ]


def create_excel_file(artefact_id: str, data: list[dict]) -> bytes:
    """Create an Excel workbook with sample data and disclaimer."""
    wb = Workbook()

    # Main data sheet
    ws = wb.active
    if ws is None:
        raise RuntimeError("No active sheet")
    ws.title = "Sample Report"

    # Write headers
    headers = list(data[0].keys())
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)

    # Write data
    for row_idx, row_data in enumerate(data, start=2):
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row_data[header])

    # Create the URL where the validity of the DID can be checked.
    check_url = f"{DIDCHECK_URL}/did/{artefact_id}"

    # Create Legal Disclaimer sheet
    disclaimer_ws = wb.create_sheet("Legal Disclaimer")
    disclaimer_ws.cell(row=1, column=1, value="Disclaimer: All data is provided AS IS...")
    disclaimer_ws.cell(row=5, column=1, value="Check the integrity of this report at:")

    # Add clickable hyperlink
    link_cell = disclaimer_ws.cell(row=6, column=1, value=check_url)
    link_cell.hyperlink = check_url
    link_cell.style = "Hyperlink"

    # Generate QR code URL with a source query param for logging.
    qr_url = f"{check_url}?source=qrcode"

    # Generate QR code for the check URL
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.ERROR_CORRECT_M,
        box_size=4,
        border=2,
    )
    qr.add_data(qr_url)
    qr.make(fit=True)
    qr_image = qr.make_image(fill_color="black", back_color="white")

    # Save QR code to bytes buffer
    qr_buffer = io.BytesIO()
    qr_image.save(qr_buffer, format="PNG")  # type: ignore[union-attr]
    qr_buffer.seek(0)

    # Add QR code to the disclaimer sheet
    img = Image(qr_buffer)
    img.anchor = "A8"  # Position below the text
    disclaimer_ws.add_image(img)

    # Save to bytes
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def calculate_multihash(data: bytes) -> str:
    """Calculate SHA-256 hash and encode as base58btc multihash.

    Returns a multihash string in base58btc format, which is the standard
    encoding used by IPFS and multiformat specifications.
    """
    # Calculate SHA-256 hash
    hash_bytes = hashlib.sha256(data).digest()

    # Create multihash: 0x12 (sha2-256) + 0x20 (32 bytes length) + hash
    multihash_bytes = bytes([0x12, 0x20]) + hash_bytes

    # Encode as base58btc (Bitcoin alphabet)
    return base58.b58encode(multihash_bytes).decode("ascii")


async def register_recommendation_with_didsvc(
    recommendation_uid: str,
    artefact_hash: str | None,
    metadata: dict,
    backlink: str,
) -> dict:
    """Register a recommendation with DIDSvc."""
    request_body = {
        "recommendation_uid": recommendation_uid,
        "artefact_hash": artefact_hash,
        "artefact_metadata": metadata,
        "backlink": backlink,
    }

    async with httpx.AsyncClient() as client:
        response = await client.post(
            f"{DIDSVC_BASE_URL}/advisory/record_recommendation",
            json=request_body,
            headers={
                "Content-Type": "application/json",
                "X-API-Token": DIDSVC_API_TOKEN,
            },
            timeout=30.0,
        )

        if response.status_code != 200:
            error_detail = response.json().get("detail", f"HTTP {response.status_code}")
            raise HTTPException(
                status_code=502,
                detail=f"DIDSvc error: {error_detail}",
            )

        return response.json()


async def register_report_with_didsvc(
    report_uid: str,
    recommendation_uid: str,
    artefact_hash: str,
    metadata: dict,
) -> dict:
    """Register a report with DIDSvc."""
    request_body = {
        "report_uid": report_uid,
        "recommendation_uid": recommendation_uid,
        "artefact_hash": artefact_hash,
        "artefact_metadata": metadata,
    }

    async with httpx.AsyncClient() as client:
        response = await client.post(
            f"{DIDSVC_BASE_URL}/advisory/record_report",
            json=request_body,
            headers={
                "Content-Type": "application/json",
                "X-API-Token": DIDSVC_API_TOKEN,
            },
            timeout=30.0,
        )

        if response.status_code != 200:
            error_detail = response.json().get("detail", f"HTTP {response.status_code}")
            raise HTTPException(
                status_code=502,
                detail=f"DIDSvc error: {error_detail}",
            )

        return response.json()


@app.get("/")
async def serve_frontend():
    """Serve the frontend HTML page."""
    return FileResponse(STATIC_DIR / "index.html", media_type="text/html")


@app.post("/generate-report")
async def generate_report():
    """
    Generate an Excel report and register it with DIDSvc.

    Returns the Excel file as a download with DID information in response headers.
    """
    # Generate unique IDs for recommendation and report
    recommendation_uid = str(uuid.uuid4())
    report_uid = str(uuid.uuid4())

    # Generate sample data and create Excel file
    data = generate_sample_data()
    excel_bytes = create_excel_file(report_uid, data)

    # Calculate hash
    artefact_hash = calculate_multihash(excel_bytes)

    # Prepare recommendation metadata
    recommendation_metadata = {
        "name": "sample_recommendation",
        "type": "recommendation",
        "description": "Recommendation for sample report",
        "generatedAt": datetime.now(timezone.utc).isoformat(),
    }

    # Create backlink URL for recommendation
    recommendation_backlink = f"{EXAMPLE_APP_URL}/recommendation/{recommendation_uid}"

    # First, register the recommendation with DIDSvc
    await register_recommendation_with_didsvc(
        recommendation_uid,
        None,  # No hash for recommendation
        recommendation_metadata,
        recommendation_backlink
    )

    # Prepare report metadata
    report_metadata = {
        "name": "sample_report",
        "type": "report",
        "format": "xlsx",
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "rowCount": len(data),
    }

    # Now register the report with DIDSvc (referencing the recommendation)
    did_response = await register_report_with_didsvc(
        report_uid,
        recommendation_uid,
        artefact_hash,
        report_metadata
    )

    # Create filename
    filename = f"report_{report_uid[:8]}.xlsx"

    # Return file as streaming response with DID info in headers
    return StreamingResponse(
        io.BytesIO(excel_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Artefact-DID": did_response["artefact_did"],
            "X-Version-DID": did_response["version_did"],
            "X-Version": str(did_response["version"]),
        },
    )


@app.get("/recommendation/{artefact_id}")
async def recommendation_details(artefact_id: str):
    """
    Recommendation details page - displays the DID and link to DIDCheck.

    This demonstrates the backlink functionality: when users click the backlink
    in DIDCheck, they come here to see the recommendation DID with a link back to DIDCheck
    for full verification details.
    """
    from string import Template

    from fastapi.responses import HTMLResponse

    # Read the template file
    template_path = STATIC_DIR / "recommendation_details.html"
    with open(template_path) as f:
        template_content = f.read()

    # Prepare the data
    did_check_url = f"{DIDCHECK_URL}/did/{artefact_id}"
    did = f"did:web:did.amd.com:{artefact_id}"

    # Render the template with actual values
    template = Template(template_content)
    html_content = template.safe_substitute(
        recommendation_uid=artefact_id,
        did=did,
        did_check_url=did_check_url
    )

    return HTMLResponse(content=html_content)


@app.post("/recommendation/{recommendation_uid}/generate-report")
async def generate_report_for_recommendation(recommendation_uid: str):
    """
    Generate a new Excel report for an existing recommendation.

    This endpoint creates a new report version referencing the existing recommendation.
    Returns the Excel file as a download with DID information in response headers.
    """
    # Generate a new unique ID for this report
    report_uid = str(uuid.uuid4())

    # Generate sample data and create Excel file
    data = generate_sample_data()
    excel_bytes = create_excel_file(report_uid, data)

    # Calculate hash
    artefact_hash = calculate_multihash(excel_bytes)

    # Prepare report metadata
    report_metadata = {
        "name": "sample_report",
        "type": "report",
        "format": "xlsx",
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "rowCount": len(data),
    }

    # Register the report with DIDSvc (referencing the existing recommendation)
    did_response = await register_report_with_didsvc(
        report_uid,
        recommendation_uid,
        artefact_hash,
        report_metadata
    )

    # Create filename
    filename = f"report_{report_uid[:8]}.xlsx"

    # Return file as streaming response with DID info in headers
    return StreamingResponse(
        io.BytesIO(excel_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Artefact-DID": did_response["artefact_did"],
            "X-Version-DID": did_response["version_did"],
            "X-Version": str(did_response["version"]),
        },
    )


@app.get("/report/{artefact_id}")
async def report_details(artefact_id: str):
    """
    Report details page - displays the report DID information.

    Note: Reports don't have backlinks in this example - only recommendations do.
    This endpoint serves as a placeholder to show report information.
    """
    from string import Template

    from fastapi.responses import HTMLResponse

    # Read the template file
    template_path = STATIC_DIR / "report_details.html"
    with open(template_path) as f:
        template_content = f.read()

    # Prepare the data
    did_check_url = f"{DIDCHECK_URL}/did/{artefact_id}"
    did = f"did:web:did.amd.com:{artefact_id}"

    # Render the template with actual values
    template = Template(template_content)
    html_content = template.substitute(did=did, did_check_url=did_check_url)

    return HTMLResponse(content=html_content)


@app.get("/health")
async def health_check():
    """Health check endpoint."""
    return {"status": "ok"}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=3001)
